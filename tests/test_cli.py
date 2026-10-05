import pytest
from openpyxl import load_workbook

from tests.conftest import FakeFetcher
from web_parser import cli


@pytest.fixture
def offline(monkeypatch, site_pages):
    """CLI работает с сохранёнными страницами вместо сети."""
    fake = FakeFetcher(site_pages)
    fake.stats = {"requests": 0, "cache_hits": 0, "retries": 0}
    monkeypatch.setattr(cli, "Fetcher", lambda **kwargs: fake)
    return fake


def test_list_categories(offline, capsys):
    assert cli.main(["--list-categories"]) == 0
    out = capsys.readouterr().out.splitlines()
    assert len(out) == 50 and "Travel" in out


def test_category_to_xlsx(offline, tmp_path, capsys):
    out = tmp_path / "travel.xlsx"
    assert cli.main(["--category", "travel", "--out", str(out)]) == 0
    assert load_workbook(out)["Товары"].max_row == 12  # шапка + 11 книг
    assert "Готово: 11 товаров" in capsys.readouterr().out


def test_two_categories_to_csv(offline, tmp_path):
    out = tmp_path / "books.csv"
    assert cli.main(["-c", "Travel", "-c", "Mystery", "-o", str(out)]) == 0
    assert len(out.read_text(encoding="utf-8-sig").splitlines()) == 1 + 11 + 32


def test_unknown_category(offline, tmp_path):
    with pytest.raises(SystemExit, match="Нет категории «Fantasyy»"):
        cli.main(["--category", "Fantasyy", "--out", str(tmp_path / "x.xlsx")])


def test_bad_extension(offline, capsys):
    assert cli.main(["--out", "result.txt"]) == 2
    assert ".xlsx или .csv" in capsys.readouterr().err


@pytest.mark.parametrize(
    ("n", "expected"),
    [(1, "1 товар"), (2, "2 товара"), (5, "5 товаров"), (11, "11 товаров"), (12, "12 товаров"),
     (21, "21 товар"), (51, "51 товар"), (104, "104 товара"), (111, "111 товаров"), (1000, "1000 товаров")],
)
def test_plural(n, expected):
    assert cli.plural(n, cli.PRODUCTS) == expected
