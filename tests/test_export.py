import csv

from openpyxl import load_workbook

from web_parser.export import category_summary, to_csv, to_xlsx
from web_parser.models import Product


def products():
    return [
        Product("Книга А", "https://s/a", "Travel", 10.0, "GBP", True, 5, 4, "upc1", "Описание", "https://s/a.jpg"),
        Product("Книга Б", "https://s/b", "Travel", 30.0, "GBP", False, None, 2, "upc2"),
        Product("Книга В", "https://s/c", "Poetry", None, "GBP", True, 1, None, "upc3"),
    ]


def test_csv_opens_in_excel(tmp_path):
    path = to_csv(products(), tmp_path / "out.csv")
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM — Excel распознает UTF-8
    rows = list(csv.reader(raw.decode("utf-8-sig").splitlines(), delimiter=";"))
    assert rows[0][0] == "Название"
    assert rows[1][:6] == ["Книга А", "https://s/a", "Travel", "10.0", "GBP", "да"]
    assert rows[2][6] == ""  # пустой остаток — пустая ячейка, а не «None»


def test_summary_by_category():
    summary = {item["category"]: item for item in category_summary(products())}
    assert summary["Travel"] == {"category": "Travel", "count": 2, "in_stock": 1, "min": 10.0, "avg": 20.0, "max": 30.0}
    assert summary["Poetry"]["avg"] is None
    assert list(summary)[0] == "Travel"  # большие категории первыми


def test_xlsx_structure(tmp_path):
    path = to_xlsx(products(), tmp_path / "out.xlsx", errors=[("https://s/x", "HTTP 503")], source="https://s/")
    wb = load_workbook(path)
    assert wb.sheetnames == ["Товары", "Сводка", "Ошибки"]
    sheet = wb["Товары"]
    assert sheet["A1"].value == "Название"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:K4"
    assert sheet["B2"].value == "открыть"
    assert sheet["B2"].hyperlink.target == "https://s/a"
    assert sheet["D2"].value == 10.0  # цена — число, по ней работают фильтры и формулы
    assert wb["Ошибки"]["B2"].value == "HTTP 503"


def test_xlsx_without_errors_and_products(tmp_path):
    wb = load_workbook(to_xlsx([], tmp_path / "empty.xlsx"))
    assert wb.sheetnames == ["Товары", "Сводка"]
